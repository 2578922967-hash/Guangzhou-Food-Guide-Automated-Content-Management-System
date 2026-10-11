"""
数据导出服务。
"""

from __future__ import annotations

import csv
from io import BytesIO
from pathlib import Path
from typing import Any

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from ..db.dao import restaurant_dao


def export_to_csv(filters: dict[str, Any] | None = None) -> bytes:
    """
    导出餐厅数据为 CSV 格式。
    
    Args:
        filters: 筛选条件
        
    Returns:
        CSV 文件的字节数据
    """
    if filters is None:
        filters = {}
    
    # 获取所有数据（不分页）
    restaurants = restaurant_dao.search(filters, limit=10000, offset=0)
    
    if not restaurants:
        return b""
    
    # 定义字段顺序
    fieldnames = [
        "shop_id",
        "name",
        "category",
        "business_district",
        "address",
        "avg_price",
        "taste_score",
        "environment_score",
        "service_score",
        "star_level",
        "reviews",
        "recommended_dish_1",
        "recommended_dish_2",
        "recommended_dish_3",
        "url",
    ]
    
    # 转换为 CSV (使用 UTF-8 BOM 以支持 Excel)
    import io
    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=fieldnames,
        extrasaction="ignore",
    )
    writer.writeheader()
    for restaurant in restaurants:
        # 确保所有值都是字符串
        row = {}
        for field in fieldnames:
            value = restaurant.get(field, "")
            row[field] = str(value) if value is not None else ""
        writer.writerow(row)
    
    # 转换为字节并添加 UTF-8 BOM（Excel 需要）
    csv_content = output.getvalue()
    return csv_content.encode('utf-8-sig')


def export_to_excel(filters: dict[str, Any] | None = None) -> bytes:
    """
    导出餐厅数据为 Excel 格式。
    
    Args:
        filters: 筛选条件
        
    Returns:
        Excel 文件的字节数据
    """
    if not HAS_OPENPYXL:
        raise ImportError("需要安装 openpyxl: pip install openpyxl")
    
    if filters is None:
        filters = {}
    
    # 获取所有数据（不分页）
    restaurants = restaurant_dao.search(filters, limit=10000, offset=0)
    
    if not restaurants:
        return b""
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "餐厅数据"
    
    # 定义字段顺序和中文标题
    fieldnames = [
        "shop_id",
        "name",
        "category",
        "business_district",
        "address",
        "avg_price",
        "taste_score",
        "environment_score",
        "service_score",
        "star_level",
        "reviews",
        "recommended_dish_1",
        "recommended_dish_2",
        "recommended_dish_3",
        "url",
    ]
    
    headers = [
        "店铺ID",
        "餐厅名称",
        "菜系",
        "商圈",
        "地址",
        "人均价格",
        "味道评分",
        "环境评分",
        "服务评分",
        "星级",
        "评论数",
        "推荐菜1",
        "推荐菜2",
        "推荐菜3",
        "链接",
    ]
    
    # 设置表头样式
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入表头
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    
    # 写入数据
    for row_idx, restaurant in enumerate(restaurants, start=2):
        for col_idx, field in enumerate(fieldnames, start=1):
            value = restaurant.get(field, "")
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[col_letter].width = adjusted_width
    
    # 保存到字节流
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output.getvalue()


def export_to_excel_pandas(filters: dict[str, Any] | None = None) -> bytes:
    """
    使用 pandas 导出为 Excel（备用方案）。
    
    Args:
        filters: 筛选条件
        
    Returns:
        Excel 文件的字节数据
    """
    if not HAS_PANDAS:
        raise ImportError("需要安装 pandas: pip install pandas")
    
    if filters is None:
        filters = {}
    
    # 获取所有数据（不分页）
    restaurants = restaurant_dao.search(filters, limit=10000, offset=0)
    
    if not restaurants:
        return b""
    
    # 转换为 DataFrame
    df = pd.DataFrame(restaurants)
    
    # 定义字段顺序和中文标题
    field_mapping = {
        "shop_id": "店铺ID",
        "name": "餐厅名称",
        "category": "菜系",
        "business_district": "商圈",
        "address": "地址",
        "avg_price": "人均价格",
        "taste_score": "味道评分",
        "environment_score": "环境评分",
        "service_score": "服务评分",
        "star_level": "星级",
        "reviews": "评论数",
        "recommended_dish_1": "推荐菜1",
        "recommended_dish_2": "推荐菜2",
        "recommended_dish_3": "推荐菜3",
        "url": "链接",
    }
    
    # 选择并重命名字段
    selected_fields = list(field_mapping.keys())
    df = df[[col for col in selected_fields if col in df.columns]]
    df = df.rename(columns=field_mapping)
    
    # 导出到 Excel
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="餐厅数据")
    
    output.seek(0)
    return output.getvalue()

